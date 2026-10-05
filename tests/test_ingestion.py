"""Unit tests for document ingestion and text chunking engine."""

import os
import shutil
import tempfile
import unittest
import docx
try:
    import pymupdf as fitz
except ImportError:
    import fitz

from src.document_processor import extract_text_from_file, chunk_text


class TestDocumentProcessor(unittest.TestCase):
    """Test suite covering text extraction, chunking, and metadata integrity."""

    def setUp(self):
        """Create a temporary test directory for sample files."""
        self.test_dir = tempfile.mkdtemp()

    def tearDown(self):
        """Clean up the temporary test directory."""
        shutil.rmtree(self.test_dir, ignore_errors=True)

    def test_text_file_extraction_and_chunking(self):
        """Test extraction and chunking of a plain .txt file."""
        sample_path = os.path.join(self.test_dir, "sample.txt")
        # Text with length > 800 chars to verify chunking and overlap
        paragraph = (
            "Natural Language Processing (NLP) is a subfield of artificial intelligence "
            "concerned with the interactions between computers and human language. "
            "Retrieval-Augmented Generation (RAG) enhances Large Language Models by "
            "fetching authoritative knowledge bases before generating a response. "
        )
        sample_text = (paragraph * 5).strip()

        with open(sample_path, "w", encoding="utf-8") as f:
            f.write(sample_text)

        # 1. Extraction assertion
        extracted = extract_text_from_file(sample_path)
        self.assertEqual(len(extracted), 1)
        self.assertEqual(extracted[0]["page"], 1)
        self.assertEqual(extracted[0]["text"], sample_text)

        # 2. Chunking assertion
        chunk_size = 500
        overlap = 100
        chunks = chunk_text(
            extracted, filename="sample.txt", chunk_size=chunk_size, overlap=overlap
        )

        self.assertGreater(len(chunks), 1)
        for i, chunk in enumerate(chunks):
            self.assertEqual(chunk["id"], f"sample.txt_p1_c{i}")
            self.assertLessEqual(len(chunk["text"]), chunk_size)
            self.assertEqual(chunk["metadata"]["filename"], "sample.txt")
            self.assertEqual(chunk["metadata"]["page"], 1)

        # Check overlap between first and second chunk
        expected_overlap_suffix = chunks[0]["text"][-overlap:]
        self.assertTrue(chunks[1]["text"].startswith(expected_overlap_suffix))

    def test_markdown_file_extraction_and_chunking(self):
        """Test extraction and chunking of a .md file."""
        sample_path = os.path.join(self.test_dir, "guide.md")
        content = "# Introduction to RAG\n\nThis is a sample markdown document for testing."
        with open(sample_path, "w", encoding="utf-8") as f:
            f.write(content)

        extracted = extract_text_from_file(sample_path)
        self.assertEqual(len(extracted), 1)
        self.assertEqual(extracted[0]["page"], 1)
        self.assertEqual(extracted[0]["text"], content)

        chunks = chunk_text(extracted, filename="guide.md", chunk_size=800, overlap=150)
        self.assertEqual(len(chunks), 1)
        self.assertEqual(chunks[0]["id"], "guide.md_p1_c0")
        self.assertEqual(chunks[0]["text"], content)
        self.assertEqual(chunks[0]["metadata"]["filename"], "guide.md")
        self.assertEqual(chunks[0]["metadata"]["page"], 1)

    def test_docx_file_extraction(self):
        """Test extraction of a .docx file using python-docx."""
        sample_path = os.path.join(self.test_dir, "sample.docx")
        doc = docx.Document()
        doc.add_paragraph("Paragraph 1: Fundamentals of RAG architectures.")
        doc.add_paragraph("Paragraph 2: Vector embedding strategies.")
        doc.save(sample_path)

        extracted = extract_text_from_file(sample_path)
        self.assertEqual(len(extracted), 1)
        self.assertEqual(extracted[0]["page"], 1)
        self.assertIn("Fundamentals of RAG architectures.", extracted[0]["text"])
        self.assertIn("Vector embedding strategies.", extracted[0]["text"])

        chunks = chunk_text(extracted, filename="sample.docx", chunk_size=800, overlap=150)
        self.assertEqual(len(chunks), 1)
        self.assertEqual(chunks[0]["id"], "sample.docx_p1_c0")
        self.assertEqual(chunks[0]["metadata"]["filename"], "sample.docx")

    def test_pdf_file_extraction(self):
        """Test extraction of a multi-page PDF using PyMuPDF (fitz)."""
        sample_path = os.path.join(self.test_dir, "sample.pdf")
        doc = fitz.open()

        page1 = doc.new_page()
        page1.insert_text((50, 50), "Content from Page 1: Overview of Generative AI.")

        page2 = doc.new_page()
        page2.insert_text((50, 50), "Content from Page 2: Vector Search with ChromaDB.")

        doc.save(sample_path)
        doc.close()

        extracted = extract_text_from_file(sample_path)
        self.assertEqual(len(extracted), 2)

        self.assertEqual(extracted[0]["page"], 1)
        self.assertIn("Overview of Generative AI", extracted[0]["text"])

        self.assertEqual(extracted[1]["page"], 2)
        self.assertIn("Vector Search with ChromaDB", extracted[1]["text"])

        chunks = chunk_text(extracted, filename="sample.pdf", chunk_size=800, overlap=150)
        self.assertEqual(len(chunks), 2)
        self.assertEqual(chunks[0]["id"], "sample.pdf_p1_c0")
        self.assertEqual(chunks[0]["metadata"]["page"], 1)
        self.assertEqual(chunks[1]["id"], "sample.pdf_p2_c0")
        self.assertEqual(chunks[1]["metadata"]["page"], 2)

    def test_error_handling_and_validation(self):
        """Test error handling for non-existent files and invalid parameters."""
        # Non-existent file
        with self.assertRaises(FileNotFoundError):
            extract_text_from_file("non_existent_file.pdf")

        # Unsupported file extension
        unsupported_path = os.path.join(self.test_dir, "data.csv")
        with open(unsupported_path, "w", encoding="utf-8") as f:
            f.write("col1,col2\n1,2")
        with self.assertRaises(ValueError):
            extract_text_from_file(unsupported_path)

        # Invalid chunking parameters
        valid_pages = [{"text": "Sample text", "page": 1}]
        with self.assertRaises(ValueError):
            chunk_text(valid_pages, "test.txt", chunk_size=0, overlap=0)
        with self.assertRaises(ValueError):
            chunk_text(valid_pages, "test.txt", chunk_size=100, overlap=100)
        with self.assertRaises(ValueError):
            chunk_text(valid_pages, "test.txt", chunk_size=100, overlap=-5)


if __name__ == "__main__":
    unittest.main()
